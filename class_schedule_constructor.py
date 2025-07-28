import sys
import ast
from datetime import datetime, timedelta
from docx import Document

def main():
    file = sys.argv[1]
    classes = []

    # Read each line as a Python dictionary using literal_eval
    with open(file, 'r', encoding='utf-8') as f:
        for line in f:
            line = line.strip()
            if line:
                cls = ast.literal_eval(line)
                classes.append(cls)

    schedule = {}

    for cls in classes:
        # Extract and format date part only
        date_str, x = cls["Day"].split(" ")
        # Convert to datetime for sorting and consistency
        day_obj = datetime.strptime(date_str, "%d/%m/%Y")
        day_formatted = day_obj.strftime("%Y/%m/%d")

        entry = f' {cls["Time"]}  | {cls["Course"]} {cls["Room"]}  {cls["Professor"]}'

        if f"{day_formatted}  {x}" not in schedule:
            schedule[f"{day_formatted}  {x}"] = []
        schedule[f"{day_formatted}  {x}"].append(entry)

    # Sort schedule by date
    sorted_schedule = dict(sorted(schedule.items()))

    for d, entries in sorted_schedule.items():

        sorted_entries = sorted(entries, key = lambda k: int(k.split("|")[0].strip()[:2]))
        sorted_schedule[d] = sorted_entries


    
    # Write to Word document
    doc = Document()
    doc.add_heading("Class Schedule", 0)

    for day, entries in sorted_schedule.items():
        doc.add_heading(day, level=1)
        for entry in entries:
            doc.add_paragraph(entry, style='List Bullet')

    doc.save("Schedule.docx")
    print("✅ Schedule written to Schedule.docx")

if __name__ == "__main__":
    main()





